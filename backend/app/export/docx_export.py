from io import BytesIO
from docx import Document
from ..schemas import ProposalDraft

def draft_to_docx(draft: ProposalDraft) -> bytes:
    doc = Document()
    doc.add_heading(draft.title or "Proposal", level=0)
    for section in draft.sections:
        doc.add_heading(section.title, level=1)
        for para in section.body.split("\n\n"):
            doc.add_paragraph(para)
        if section.citations:
            doc.add_paragraph("Sources: " + ", ".join(f"[{c}]" for c in section.citations)).italic = True
    if draft.pricing:
        doc.add_heading("Pricing", level=1)
        table = doc.add_table(rows=1, cols=4)
        for cell, head in zip(table.rows[0].cells, ["Item", "Amount", "Basis", "Note"]):
            cell.text = head
        for line in draft.pricing:
            row = table.add_row().cells
            row[0].text, row[1].text = line.item, f"${line.amount:,.0f}"
            row[2].text, row[3].text = line.basis, line.note
    buf = BytesIO(); doc.save(buf); return buf.getvalue()
